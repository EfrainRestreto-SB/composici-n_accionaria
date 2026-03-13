import openpyxl
import pandas as pd

print("=" * 100)
print("COMPARACIÓN DETALLADA: data.xlsx vs prueba.xlsx")
print("=" * 100)

# Leer ambos archivos
wb_data = openpyxl.load_workbook('data.xlsx')
wb_prueba = openpyxl.load_workbook('prueba.xlsx')

ws_data = wb_data.active
ws_prueba = wb_prueba.active

print("\nPRIMERAS 45 FILAS - COMPARACIÓN LADO A LADO")
print("-" * 100)
print(f"{'FILA':<6}{'DATA.XLSX (Col A)':<40}{'DATA.XLSX (Col B)':<12}{'PRUEBA.XLSX (Col A)':<40}{'PRUEBA.XLSX (Col B)':<12}")
print("-" * 100)

for i in range(45):
    row_num = i + 1
    
    # Data.xlsx
    data_a = ws_data.cell(row_num, 1).value
    data_b = ws_data.cell(row_num, 2).value
    
    # Prueba.xlsx  
    prueba_a = ws_prueba.cell(row_num, 1).value
    prueba_b = ws_prueba.cell(row_num, 2).value
    
    # Solo mostrar filas con contenido
    if data_a or data_b or prueba_a or prueba_b:
        data_a_str = str(data_a) if data_a is not None else ""
        data_b_str = str(data_b) if data_b is not None else ""
        prueba_a_str = str(prueba_a) if prueba_a is not None else ""
        prueba_b_str = str(prueba_b) if prueba_b is not None else ""
        
        # Truncar nombres largos
        data_a_str = data_a_str[:38] if len(data_a_str) > 38 else data_a_str
        data_b_str = data_b_str[:10] if len(data_b_str) > 10 else data_b_str
        prueba_a_str = prueba_a_str[:38] if len(prueba_a_str) > 38 else prueba_a_str
        prueba_b_str = prueba_b_str[:10] if len(prueba_b_str) > 10 else prueba_b_str
        
        print(f"{i:<6}{data_a_str:<40}{data_b_str:<12}{prueba_a_str:<40}{prueba_b_str:<12}")

print("\n" + "=" * 100)
print("CONCLUSIÓN:")
print("¿Son exactamente iguales en estructura? Verificar columnas A y B")
print("=" * 100)
