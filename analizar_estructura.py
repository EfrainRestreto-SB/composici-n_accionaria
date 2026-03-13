import openpyxl

wb = openpyxl.load_workbook('prueba.xlsx')
ws = wb.active

print('=== ESTRUCTURA DETALLADA PRUEBA.XLSX (filas 0-40) ===')
print()

for i in range(40):
    row_num = i + 1
    a = ws.cell(row_num, 1).value
    b = ws.cell(row_num, 2).value
    c = ws.cell(row_num, 3).value
    d = ws.cell(row_num, 4).value
    
    if a or b or c or d:
        a_str = str(a) if a is not None else ""
        b_str = str(b) if b is not None else ""
        c_str = str(c) if c is not None else ""
        d_str = str(d) if d is not None else ""
        print(f"F{i:2d}: [{a_str:35s}] | [{b_str:10s}] | [{c_str:10s}] | [{d_str:10s}]")
