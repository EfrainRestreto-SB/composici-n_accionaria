#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Parser Genérico de Formato Jerárquico a Formato Relacional
Convierte reportes con estructura jerárquica a formato tabla relacional
que el sistema Java puede procesar.

Formato de entrada (jerárquico):
| Entidad (A)        | % Directo (B) | % Acumulado (C) |
|--------------------|---------------|-----------------|
| ROOT ENTITY        | 1.0           | 1.0             |
| SHAREHOLDER 1      | 0.6           | =B*C            |
| SHAREHOLDER 2      | 0.4           | =B*C            |
| SHAREHOLDER 1      | (vacío)       | 0.6             | <- Inicia desglose
| SUB-SHAREHOLDER A  | 0.5           | =B*C            |
| SUB-SHAREHOLDER B  | 0.5           | =B*C            |

Formato de salida (relacional):
| Entidad (A)   | Accionista (B)    | % (C) |
|---------------|-------------------|-------|
| ROOT ENTITY   | SHAREHOLDER 1     | 60    |
| ROOT ENTITY   | SHAREHOLDER 2     | 40    |
| SHAREHOLDER 1 | SUB-SHAREHOLDER A | 50    |
| SHAREHOLDER 1 | SUB-SHAREHOLDER B | 50    |
"""

import sys
import io
import openpyxl
from pathlib import Path
from collections import defaultdict

# Configurar codificación
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

class HierarchicalParser:
    """Parser para convertir formato jerárquico a relacional"""
    
    def __init__(self, excel_path):
        self.excel_path = excel_path
        self.wb = None
        self.ws = None
        self.relationships = []  # Lista de (entidad, accionista, porcentaje)
        self.root_entity = None
        
    def parse(self):
        """Parsea el archivo Excel jerárquico"""
        print(f"=== PARSER GENÉRICO DE FORMATO JERÁRQUICO ===")
        print(f"Archivo: {self.excel_path}")
        print()
        
        self.wb = openpyxl.load_workbook(self.excel_path)
        self.ws = self.wb.active
        
        # Paso 1: Detectar entidad raíz
        self.root_entity = self._detect_root_entity()
        if not self.root_entity:
            raise ValueError("No se pudo detectar la entidad raíz en el archivo")
        
        print(f"✓ Entidad raíz detectada: {self.root_entity}")
        
        # Paso 2: Parsear jerarquía
        self._parse_hierarchy()
        
        print(f"✓ Relaciones extraídas: {len(self.relationships)}")
        
        return self.relationships
    
    def _detect_root_entity(self):
        """
        Detecta la entidad raíz buscando la primera fila donde:
        - Columna A tiene texto
        - Columna B tiene valor 1.0 o cercano
        - Es la primera entidad que aparece
        """
        for row_idx in range(1, min(20, self.ws.max_row + 1)):
            cell_a = self.ws.cell(row_idx, 1).value
            cell_b = self.ws.cell(row_idx, 2).value
            
            if cell_a and isinstance(cell_a, str) and cell_a.strip():
                # Si columna B es 1.0, probablemente es la raíz
                if cell_b is not None:
                    try:
                        value_b = float(cell_b)
                        if abs(value_b - 1.0) < 0.01:
                            return cell_a.strip()
                    except (ValueError, TypeError):
                        pass
                
                # Si es la primera entidad con nombre válido, asumirla como raíz
                if self._is_entity_name(cell_a):
                    return cell_a.strip()
        
        return None
    
    def _is_entity_name(self, text):
        """Determina si un texto parece ser nombre de entidad válido"""
        if not text or not isinstance(text, str):
            return False
        
        text = text.strip()
        if len(text) < 3:
            return False
        
        # Debe contener al menos una letra
        if not any(c.isalpha() for c in text):
            return False
        
        return True
    
    def _parse_hierarchy(self):
        """
        Parsea la jerarquía completa del archivo.
        
        Lógica:
        1. Cuando columna A tiene entidad Y columna B tiene porcentaje → es una relación directa
        2. El "padre" es la última entidad con columna C (% acumulado) encontrada antes
        3. Cuando columna A repite entidad con columna B vacía → inicia nuevo nivel
        """
        current_parent = None
        entity_stack = []  # Stack para tracking de contexto jerárquico
        
        for row_idx in range(1, self.ws.max_row + 1):
            cell_a = self.ws.cell(row_idx, 1).value
            cell_b = self.ws.cell(row_idx, 2).value
            cell_c = self.ws.cell(row_idx, 3).value
            
            # Saltar filas vacías o de encabezado
            if not cell_a or not isinstance(cell_a, str):
                continue
            
            entity_name = cell_a.strip()
            
            # Saltar si parece ser encabezado
            if self._is_header_row(entity_name):
                continue
            
            # Caso 1: Entidad con porcentaje directo (columna B)
            # → Es un accionista del padre actual
            if cell_b is not None:
                try:
                    percentage = self._parse_percentage(cell_b)
                    if percentage > 0:
                        # Determinar el padre
                        parent = current_parent if current_parent else self.root_entity
                        
                        # Evitar relaciones autoreferenciales (entidad propiedad de sí misma)
                        if parent != entity_name:
                            # Agregar relación
                            self.relationships.append((parent, entity_name, percentage))
                            print(f"  {parent} ← {entity_name} ({percentage:.2f}%)")
                        else:
                            print(f"  [Ignorado] Autorreferencia: {entity_name}")
                        
                except (ValueError, TypeError):
                    pass
            
            # Caso 2: Entidad sin porcentaje directo pero con porcentaje acumulado (columna C)
            # → Se convierte en el nuevo padre para las siguientes filas
            elif cell_c is not None:
                try:
                    accumulated = self._parse_percentage_or_formula(cell_c)
                    if accumulated > 0:
                        current_parent = entity_name
                        print(f"\n→ Nuevo contexto: {entity_name} ({accumulated:.2f}%)")
                except (ValueError, TypeError):
                    pass
    
    def _is_header_row(self, text):
        """Detecta si una fila es encabezado"""
        text_upper = text.upper()
        header_keywords = ['DESGLOSE', 'COMPOSICION', 'ACCIONARIA', 'ENTIDAD', 
                          'ACCIONISTA', 'PARTICIPACION', 'PORCENTAJE']
        return any(keyword in text_upper for keyword in header_keywords)
    
    def _parse_percentage(self, value):
        """Convierte valor a porcentaje (0-100)"""
        if value is None:
            return 0.0
        
        if isinstance(value, str):
            value = value.strip()
            if not value or value == 'nan':
                return 0.0
        
        num = float(value)
        
        # Si está entre 0-1, convertir a porcentaje
        if 0 <= num <= 1:
            return num * 100
        
        # Si está entre 0-100, ya es porcentaje
        if 0 <= num <= 100:
            return num
        
        return 0.0
    
    def _parse_percentage_or_formula(self, value):
        """Parsea porcentaje o fórmula de Excel"""
        if value is None:
            return 0.0
        
        # Si es string que parece fórmula, no podemos evaluarla aquí
        if isinstance(value, str) and value.strip().startswith('='):
            # Intentar extraer valor evaluado de la celda
            return 0.0
        
        return self._parse_percentage(value)
    
    def save_to_excel(self, output_path):
        """Guarda las relaciones en formato relacional"""
        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = "Datos"
        
        # Header
        ws_out['A1'] = 'Entidad'
        ws_out['B1'] = 'Accionista'
        ws_out['C1'] = '% Participación'
        
        # Aplicar formato
        from openpyxl.styles import Font, PatternFill
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
        for cell in [ws_out['A1'], ws_out['B1'], ws_out['C1']]:
            cell.font = header_font
            cell.fill = header_fill
        
        # Datos
        for idx, (entity, shareholder, percentage) in enumerate(self.relationships, start=2):
            ws_out[f'A{idx}'] = entity
            ws_out[f'B{idx}'] = shareholder
            ws_out[f'C{idx}'] = percentage
        
        wb_out.save(output_path)
        print(f"\n✓ Archivo generado: {output_path}")
        print(f"  Relaciones: {len(self.relationships)}")
        
        # Estadísticas
        self._print_statistics()
    
    def _print_statistics(self):
        """Imprime estadísticas de la conversión"""
        print(f"\n=== ESTADÍSTICAS ===")
        print(f"Entidad raíz: {self.root_entity}")
        print(f"Total relaciones: {len(self.relationships)}")
        
        # Contar entidades únicas
        entities = set()
        shareholders = set()
        for entity, shareholder, _ in self.relationships:
            entities.add(entity)
            shareholders.add(shareholder)
        
        print(f"Entidades únicas: {len(entities)}")
        print(f"Accionistas únicos: {len(shareholders)}")
        
        # Validar suma de porcentajes por entidad
        entity_totals = defaultdict(float)
        for entity, _, percentage in self.relationships:
            entity_totals[entity] += percentage
        
        print(f"\nValidación de porcentajes:")
        for entity, total in sorted(entity_totals.items()):
            status = "✓" if abs(total - 100) < 1 else "⚠"
            print(f"  {status} {entity}: {total:.2f}%")


def main():
    """Función principal"""
    if len(sys.argv) < 2:
        print("[ERROR] Uso: python parse_hierarchical_format.py <archivo_excel>")
        return False
    
    excel_file = sys.argv[1]
    
    if not Path(excel_file).exists():
        print(f"[ERROR] Archivo no encontrado: {excel_file}")
        return False
    
    try:
        # Parsear archivo
        parser = HierarchicalParser(excel_file)
        parser.parse()
        
        # Generar archivo de salida
        input_path = Path(excel_file)
        output_file = input_path.parent / f"{input_path.stem}_converted.xlsx"
        parser.save_to_excel(str(output_file))
        
        print(f"\n[OK] Conversión completada exitosamente")
        return True
        
    except Exception as e:
        print(f"\n[ERROR] Error durante la conversión: {e}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
